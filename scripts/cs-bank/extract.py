# -*- coding: utf-8 -*-
"""把本地备考资料抽取为 JSONL（每文件一条：id/root/relpath/ext/chars/text）。

用法：py extract.py <out.jsonl>
"""
import json
import os
import sys
import zipfile
import re

ROOTS = [
    ("grade", r"F:\All_Files\计算机等级考试资料"),
    ("kaoyan_real", r"F:\All_Files\408考研资料\考研真题\cskaoyan-master"),
    ("kaoyan_mock", r"F:\All_Files\408考研资料\模拟题\7.20"),
    ("interview", r"F:\All_Files\book-master\面试"),
    ("book", r"F:\All_Files\book-master\Book"),
]

MAX_CHARS = 3_000_000


def safe_id(root_key, relpath):
    s = root_key + "_" + relpath.replace("\\", "/")
    s = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", s)
    return s[:180]


def read_pdf(path):
    import fitz
    doc = fitz.open(path)
    parts = []
    total = 0
    for i in range(doc.page_count):
        try:
            t = doc[i].get_text()
        except Exception:
            t = ""
        parts.append(t)
        total += len(t)
        if total > MAX_CHARS:
            break
    doc.close()
    return "\n".join(parts)[:MAX_CHARS]


def read_docx(path):
    from docx import Document
    d = Document(path)
    return "\n".join(p.text for p in d.paragraphs)[:MAX_CHARS]


def read_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)[:MAX_CHARS]


def read_text(path):
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()[:MAX_CHARS]
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(path, "rb") as f:
        return f.read()[:MAX_CHARS].decode("utf-8", "ignore")


def read_zip_text(path):
    """从 zip 里拿出所有 .txt/.md 文本（用于等级考试资料里的压缩包）。"""
    out = []
    try:
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if name.lower().endswith((".txt", ".md", ".csv")):
                    try:
                        out.append(z.read(name).decode("utf-8", "ignore"))
                    except Exception:
                        pass
                if sum(len(x) for x in out) > MAX_CHARS:
                    break
    except Exception:
        return ""
    return "\n".join(out)[:MAX_CHARS]


HANDLERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".pptx": read_pptx,
    ".txt": read_text,
    ".md": read_text,
    ".zip": read_zip_text,
}
SKIP_EXT = {".rar", ".gz", ".doc", ".ppt", ".xlsx", ".png", ".jpg", ".jpeg", ".msi", ".exe", ".py", ".gitignore"}


def main():
    out_path = sys.argv[1]
    manifest = []
    written = 0
    with open(out_path, "w", encoding="utf-8") as out:
        for root_key, root in ROOTS:
            if not os.path.isdir(root):
                print("MISSING ROOT", root)
                continue
            for dirpath, _dirs, files in os.walk(root):
                for fn in files:
                    full = os.path.join(dirpath, fn)
                    ext = os.path.splitext(fn)[1].lower()
                    rel = os.path.relpath(full, root)
                    handler = HANDLERS.get(ext)
                    if handler is None:
                        manifest.append({"root": root_key, "relpath": rel, "ext": ext, "status": "skip"})
                        continue
                    try:
                        text = handler(full)
                    except Exception as e:  # noqa
                        manifest.append({"root": root_key, "relpath": rel, "ext": ext, "status": "error", "error": str(e)[:200]})
                        continue
                    text = text.strip()
                    if len(text) < 40:
                        manifest.append({"root": root_key, "relpath": rel, "ext": ext, "status": "empty", "chars": len(text)})
                        continue
                    rec = {
                        "id": safe_id(root_key, rel),
                        "root": root_key,
                        "relpath": rel.replace("\\", "/"),
                        "ext": ext,
                        "chars": len(text),
                        "text": text,
                    }
                    out.write(json.dumps(rec, ensure_ascii=False) + "\n")
                    written += 1
                    manifest.append({"root": root_key, "relpath": rel, "ext": ext, "status": "ok", "chars": len(text)})
    with open(out_path + ".manifest.json", "w", encoding="utf-8") as mf:
        json.dump(manifest, mf, ensure_ascii=False, indent=1)
    ok = sum(1 for m in manifest if m["status"] == "ok")
    print(f"files written: {written} (ok={ok}, total={len(manifest)})")


if __name__ == "__main__":
    main()
